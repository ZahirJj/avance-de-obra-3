"""Datos sintéticos y gráficos de seguimiento de obra, basados en un Excel editable."""
from pathlib import Path
import argparse
import os

BASE = Path(__file__).resolve().parent
# Mantiene la caché de Matplotlib dentro del proyecto.
CACHE = BASE.parent / "work" / "matplotlib"
CACHE.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(CACHE))

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter, FuncFormatter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NEGRO = "#090710"
MORADO = "#A978FF"
CLARO = "#D8BFFF"
GRIS = "#A49BB5"


def generar_excel(ruta: Path) -> None:
    """Genera 32 semanas, 12 partidas y cantidades reales hasta el corte."""
    rng = np.random.default_rng(42)
    # ID, partida, unidad, metrado contractual, precio unitario S/, inicio, fin, retraso, productividad.
    definiciones = [
        ("P01", "Obras provisionales", "glb", 1, 180000, 1, 3, 0, 1.00),
        ("P02", "Movimiento de tierras", "m³", 4200, 85, 2, 7, 1, 0.96),
        ("P03", "Cimentaciones", "m³", 950, 780, 5, 11, 1, 0.98),
        ("P04", "Estructuras de concreto", "m³", 2600, 1150, 9, 21, 2, 0.94),
        ("P05", "Albañilería", "m²", 6800, 125, 15, 25, 2, 0.96),
        ("P06", "Instalaciones sanitarias", "pto", 280, 980, 16, 28, 2, 0.94),
        ("P07", "Instalaciones eléctricas", "pto", 420, 820, 17, 29, 1, 1.03),
        ("P08", "Tarrajeos y revestimientos", "m²", 11000, 48, 19, 29, 2, 0.95),
        ("P09", "Pisos y enchapes", "m²", 3800, 145, 23, 30, 2, 0.97),
        ("P10", "Carpintería y cerrajería", "und", 180, 1750, 24, 31, 1, 0.98),
        ("P11", "Pintura", "m²", 12500, 24, 26, 32, 1, 1.00),
        ("P12", "Pruebas y entrega", "glb", 1, 95000, 30, 32, 1, 1.00),
    ]
    partidas = pd.DataFrame(definiciones, columns=[
        "id", "partida", "unidad", "cantidad_total", "precio_unitario",
        "semana_inicio", "semana_fin", "retraso_semanas", "factor_productividad",
    ])
    partidas["presupuesto"] = partidas.cantidad_total * partidas.precio_unitario
    partidas["peso_presupuesto"] = partidas.presupuesto / partidas.presupuesto.sum()
    semanas, corte = 32, 22
    inicio = pd.Timestamp("2026-01-05")
    filas = []
    for p in partidas.itertuples(index=False):
        duracion = p.semana_fin - p.semana_inicio + 1
        perfil = np.sin(np.pi * (np.arange(duracion) + 0.5) / duracion) ** 1.6
        plan = p.cantidad_total * perfil / perfil.sum()
        # Una productividad menor prolonga la ejecución. El perfil completo suma
        # el metrado contractual, evitando remanentes que nunca se ejecutarían.
        duracion_real = max(1, int(np.ceil(duracion / p.factor_productividad)))
        perfil_real = np.sin(np.pi * (np.arange(duracion_real) + 0.5) / duracion_real) ** 1.6
        perfil_real *= rng.uniform(0.93, 1.07, duracion_real)
        propuesta = p.cantidad_total * perfil_real / perfil_real.sum()
        for semana in range(1, semanas + 1):
            ip = semana - p.semana_inicio
            ir = semana - p.semana_inicio - p.retraso_semanas
            qp = float(plan[ip]) if 0 <= ip < duracion else 0.0
            qr = float(propuesta[ir]) if semana <= corte and 0 <= ir < duracion_real else (0.0 if semana <= corte else np.nan)
            filas.append((semana, inicio + pd.Timedelta(weeks=semana - 1), p.id, qp, qr))
    seguimiento = pd.DataFrame(filas, columns=["semana", "fecha_inicio", "id", "cantidad_plan", "cantidad_real"])
    parametros = pd.DataFrame({"parametro": ["proyecto", "moneda", "fecha_inicio", "semanas", "semana_corte", "semilla"],
        "valor": ["Edificio residencial de 5 niveles — ejemplo sintético", "PEN (S/)", inicio.strftime("%Y-%m-%d"), semanas, corte, 42]})
    notas = pd.DataFrame({"criterio": [
        "Datos exclusivamente sintéticos para un ejercicio de control de obra.",
        "Una fila por semana y partida: 32 × 12 = 384 registros.",
        "Metrados planificados distribuidos con un perfil suave y actividades superpuestas.",
        "El avance global se pondera por presupuesto; no se suman unidades heterogéneas.",
        "Cantidad real vacía después del corte: aún no observada, no equivale a cero.",
        "Valorización = metrado ejecutado × precio contractual; no representa gasto real.",
        "El programa admite editar metrados semanales y semana de corte en el Excel.",
    ]})
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        for nombre, tabla in [("Parametros", parametros), ("Partidas", partidas), ("Seguimiento", seguimiento), ("Notas", notas)]:
            tabla.to_excel(writer, sheet_name=nombre, index=False)
        for hoja in writer.book.worksheets:
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions
            for celda in hoja[1]:
                celda.fill = PatternFill("solid", fgColor="643A99")
                celda.font = Font(color="FFFFFF", bold=True)
                celda.alignment = Alignment(vertical="center")
            for columna in hoja.columns:
                letra = get_column_letter(columna[0].column)
                hoja.column_dimensions[letra].width = min(85, max(14, max(len(str(c.value or "")) for c in columna) + 2))
                for celda in list(columna)[1:]:
                    if isinstance(celda.value, (float, int)):
                        celda.number_format = "#,##0.00"


def leer_y_calcular(ruta: Path):
    """Lee el Excel, valida cantidades y calcula indicadores ponderados."""
    parametros = pd.read_excel(ruta, sheet_name="Parametros").set_index("parametro")["valor"]
    partidas = pd.read_excel(ruta, sheet_name="Partidas")
    seguimiento = pd.read_excel(ruta, sheet_name="Seguimiento")
    corte, semanas = int(parametros["semana_corte"]), int(parametros["semanas"])
    if not 1 <= corte <= semanas:
        raise ValueError("La semana de corte debe estar dentro del plazo de obra.")
    if partidas.id.duplicated().any() or seguimiento.duplicated(["semana", "id"]).any():
        raise ValueError("Hay IDs o registros semana/partida duplicados.")
    if (partidas[["cantidad_total", "precio_unitario"]] <= 0).any().any() or partidas[["cantidad_total", "precio_unitario"]].isna().any().any():
        raise ValueError("Metrados y precios contractuales deben ser positivos.")
    esperado = pd.MultiIndex.from_product([range(1, semanas + 1), partidas.id], names=["semana", "id"])
    observado = pd.MultiIndex.from_frame(seguimiento[["semana", "id"]])
    if len(esperado.difference(observado)) or len(observado.difference(esperado)):
        raise ValueError("Debe existir una fila por cada semana y partida, sin IDs desconocidos.")
    pasado = seguimiento.semana <= corte
    if seguimiento.cantidad_plan.isna().any() or seguimiento.loc[pasado, "cantidad_real"].isna().any():
        raise ValueError("Faltan metrados planificados o reales hasta la fecha de corte.")
    if (seguimiento[["cantidad_plan", "cantidad_real"]] < 0).any().any():
        raise ValueError("Los metrados no pueden ser negativos.")
    if seguimiento.loc[~pasado, "cantidad_real"].notna().any():
        raise ValueError("Deje vacío el avance real posterior a la semana de corte.")
    totales = seguimiento.groupby("id").cantidad_plan.sum().reindex(partidas.id).to_numpy()
    if not np.allclose(totales, partidas.cantidad_total, rtol=1e-6, atol=1e-6):
        raise ValueError("La suma del metrado planificado debe coincidir con el contractual de cada partida.")
    partidas["presupuesto"] = partidas.cantidad_total * partidas.precio_unitario
    presupuesto = partidas.presupuesto.sum()
    detalle = seguimiento.merge(partidas[["id", "precio_unitario"]], on="id", validate="many_to_one")
    detalle["valor_plan"] = detalle.cantidad_plan * detalle.precio_unitario
    detalle["valor_real"] = detalle.cantidad_real * detalle.precio_unitario
    semanal = detalle.groupby("semana")[["valor_plan", "valor_real"]].sum(min_count=1)
    semanal["plan_acumulado"] = semanal.valor_plan.cumsum()
    semanal["real_acumulado"] = semanal.valor_real.cumsum()
    semanal["plan_pct"] = semanal.plan_acumulado / presupuesto * 100
    semanal["real_pct"] = semanal.real_acumulado / presupuesto * 100
    acumulado = seguimiento.loc[pasado].groupby("id")[["cantidad_plan", "cantidad_real"]].sum()
    resumen = partidas.merge(acumulado, on="id", validate="one_to_one")
    if (resumen.cantidad_real > resumen.cantidad_total + 1e-6).any():
        raise ValueError("El metrado real supera el contractual de una partida.")
    resumen["plan_pct"] = resumen.cantidad_plan / resumen.cantidad_total * 100
    resumen["real_pct"] = resumen.cantidad_real / resumen.cantidad_total * 100
    resumen["brecha_pp"] = resumen.real_pct - resumen.plan_pct
    return parametros, corte, presupuesto, semanal, resumen


def graficar(parametros, corte, presupuesto, semanal, resumen, salida: Path):
    salida.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"figure.facecolor": NEGRO, "axes.facecolor": NEGRO,
        "savefig.facecolor": NEGRO, "text.color": "#F5EDFF", "axes.labelcolor": CLARO,
        "xtick.color": GRIS, "ytick.color": GRIS, "axes.edgecolor": "#40334F",
        "font.family": "DejaVu Sans", "font.size": 11, "axes.titleweight": "bold"})
    fila = semanal.loc[corte]
    fig, ax = plt.subplots(figsize=(13, 7), layout="constrained")
    x = np.r_[0, semanal.index.to_numpy()]
    ax.plot(x, np.r_[0, semanal.plan_pct], color=CLARO, linestyle="--", linewidth=2.5, label="Avance planificado")
    observado = semanal.loc[:corte]
    xr, yr = np.r_[0, observado.index], np.r_[0, observado.real_pct]
    ax.plot(xr, yr, color=MORADO, linewidth=3, marker="o", markersize=4, label="Avance real al corte")
    ax.fill_between(xr, yr, color=MORADO, alpha=0.12)
    ax.axvline(corte, color=GRIS, linestyle=":", alpha=0.8, label=f"Corte: semana {corte}")
    ax.scatter([corte, corte], [fila.plan_pct, fila.real_pct], color=[CLARO, MORADO], s=65, zorder=4)
    ax.set(title="Curva S · avance acumulado de obra", xlabel="Semana de ejecución", ylabel="Avance ponderado por presupuesto", xlim=(0, len(semanal) + 0.5), ylim=(0, 105))
    ax.yaxis.set_major_formatter(PercentFormatter(100))
    ax.grid(color="#4B3A61", alpha=0.3)
    ax.legend(loc="upper left", facecolor="#171021", edgecolor="#40334F", labelcolor=CLARO)
    texto = f"Plan: {fila.plan_pct:.1f}%  |  Real: {fila.real_pct:.1f}%  |  Brecha: {fila.real_pct-fila.plan_pct:+.1f} pp\nPresupuesto contractual: S/ {presupuesto:,.0f} · datos sintéticos"
    fig.text(0.5, 0.01, texto, ha="center", color=CLARO, fontsize=10)
    fig.get_layout_engine().set(rect=(0, 0.09, 1, 0.91))
    guardar(fig, salida / "curva_s")
    fig, ax = plt.subplots(figsize=(14, 9), layout="constrained")
    y = np.arange(len(resumen))
    ax.barh(y - 0.18, resumen.plan_pct, height=0.32, color=CLARO, alpha=0.75, label="Planificado al corte")
    ax.barh(y + 0.18, resumen.real_pct, height=0.32, color=MORADO, label="Ejecutado al corte")
    for i, p in enumerate(resumen.itertuples(index=False)):
        ax.text(p.plan_pct + 1, i - 0.18, f"{p.plan_pct:.1f}%", va="center", color=CLARO, fontsize=9)
        ax.text(p.real_pct + 1, i + 0.18, f"{p.real_pct:.1f}%", va="center", color=MORADO, fontsize=9)
    ax.set_yticks(y, resumen.partida)
    ax.invert_yaxis()
    ax.set(title=f"Avance de obra por partidas · semana {corte}", xlabel="Metrado acumulado / metrado contractual", xlim=(0, 116))
    ax.xaxis.set_major_formatter(PercentFormatter(100))
    ax.grid(axis="x", color="#4B3A61", alpha=0.3)
    ax.set_axisbelow(True)
    ax.legend(loc="lower right", facecolor="#171021", edgecolor="#40334F", labelcolor=CLARO)
    guardar(fig, salida / "avance_partidas")
    fig, ax = plt.subplots(figsize=(13, 6), layout="constrained")
    ax.bar(semanal.index - 0.18, semanal.valor_plan, width=0.36, color=CLARO, alpha=0.65, label="Planificado")
    ax.bar(observado.index + 0.18, observado.valor_real, width=0.36, color=MORADO, label="Ejecutado al corte")
    ax.axvline(corte + 0.5, color=GRIS, linestyle=":")
    ax.yaxis.set_major_formatter(FuncFormatter(lambda valor, _: f"{valor/1000:,.0f}"))
    ax.set(title="Valorización semanal a precios contractuales", xlabel="Semana de ejecución", ylabel="Miles de soles (S/)")
    ax.grid(axis="y", color="#4B3A61", alpha=0.3)
    ax.set_axisbelow(True)
    ax.legend(facecolor="#171021", edgecolor="#40334F", labelcolor=CLARO)
    guardar(fig, salida / "valorizacion_semanal")
    return fila


def guardar(fig, base):
    fig.savefig(base.with_suffix(".png"), dpi=180)
    fig.savefig(base.with_suffix(".pdf"))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", type=Path, default=BASE / "datos_obra.xlsx")
    parser.add_argument("--salida", type=Path, default=BASE)
    parser.add_argument("--regenerar", action="store_true", help="Sobrescribe el Excel con datos sintéticos reproducibles.")
    parser.add_argument("--mostrar", action="store_true", help="Abre los gráficos en ventanas interactivas.")
    args = parser.parse_args()
    if args.regenerar or not args.excel.exists():
        generar_excel(args.excel)
    parametros, corte, presupuesto, semanal, resumen = leer_y_calcular(args.excel)
    fila = graficar(parametros, corte, presupuesto, semanal, resumen, args.salida)
    with pd.ExcelWriter(args.salida / "resumen_avance.xlsx", engine="openpyxl") as writer:
        semanal.to_excel(writer, sheet_name="Curva_S")
        resumen.to_excel(writer, sheet_name="Avance_partidas", index=False)
    print(f"Presupuesto: S/ {presupuesto:,.2f}")
    print(f"Semana {corte}: plan {fila.plan_pct:.2f}% | real {fila.real_pct:.2f}% | brecha {fila.real_pct-fila.plan_pct:+.2f} pp")
    print(f"Gráficos y resumen guardados en: {args.salida.resolve()}")
    if args.mostrar:
        plt.show()
    else:
        plt.close("all")


if __name__ == "__main__":
    main()
